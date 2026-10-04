"""File parsers. Everything becomes Markdown with headings (for breadcrumbs) and page markers.

Routing: PDFs go to LlamaParse when they look complex or scanned (and a key is configured), otherwise
PyMuPDF. Office, spreadsheets, databases, e-mail and chat exports are parsed locally.
"""

from __future__ import annotations

import csv
import email
import email.policy
import html
import json
import logging
import mailbox
import re
import sqlite3
from dataclasses import dataclass, field
from datetime import date, datetime
from email.utils import parsedate_to_datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from ..config import Settings
from .normalize_fa import detect_lang, normalize_persian

log = logging.getLogger(__name__)

TEXT_EXTS = {".txt", ".md", ".markdown", ".text", ".log"}
MAX_SHEET_ROWS = 5000
MAX_DB_ROWS = 500


class UnsupportedFormat(Exception):
    pass


@dataclass
class ParsedDocument:
    markdown: str
    title: str | None = None
    source_type: str = "upload"
    lang: str | None = None
    doc_date: date | None = None
    page_count: int | None = None
    parser: str = "text"
    metadata: dict[str, Any] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------


def _md_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [
        [(c or "").replace("|", "\\|").replace("\n", " ").strip() for c in r] + [""] * (width - len(r))
        for r in rows
    ]
    header = norm[0]
    out = ["| " + " | ".join(header) + " |", "| " + " | ".join("---" for _ in header) + " |"]
    out += ["| " + " | ".join(r) + " |" for r in norm[1:]]
    return "\n".join(out)


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _title_from_text(text: str) -> str | None:
    for line in text.splitlines():
        s = line.strip().lstrip("#").strip()
        if 3 <= len(s) <= 120:
            return s
    return None


class _HTMLToText(HTMLParser):
    BLOCK = {"p", "div", "br", "li", "tr", "section", "article", "blockquote", "pre"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: Any) -> None:
        if tag in ("script", "style", "noscript"):
            self._skip += 1
        elif tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self.parts.append("\n\n" + "#" * int(tag[1]) + " ")
        elif tag in self.BLOCK:
            self.parts.append("\n")
        elif tag == "td" or tag == "th":
            self.parts.append(" | ")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style", "noscript") and self._skip:
            self._skip -= 1
        elif tag in ("h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "tr"):
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data)


def html_to_text(raw: str) -> str:
    p = _HTMLToText()
    p.feed(raw)
    text = html.unescape("".join(p.parts))
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def _pdf_page_markdown(page: Any) -> tuple[str, int]:
    """Markdown for one page using font-size heuristics for headings; returns (text, table_count)."""
    tables_md: list[str] = []
    table_boxes: list[Any] = []
    try:
        found = page.find_tables()
        for t in found.tables:
            try:
                md = t.to_markdown()
                if md and md.count("|") > 4:
                    tables_md.append(md.strip())
                    table_boxes.append(t.bbox)
            except Exception:  # noqa: BLE001
                continue
    except Exception:  # noqa: BLE001
        pass

    lines_out: list[str] = []
    try:
        d = page.get_text("dict")
        sizes: list[float] = []
        for b in d.get("blocks", []):
            for ln in b.get("lines", []):
                for sp in ln.get("spans", []):
                    if sp.get("text", "").strip():
                        sizes.append(float(sp.get("size", 0)))
        sizes.sort()
        median = sizes[len(sizes) // 2] if sizes else 0.0
        for b in d.get("blocks", []):
            if b.get("type", 0) != 0:
                continue
            bbox = b.get("bbox")
            if bbox and any(_inside(bbox, tb) for tb in table_boxes):
                continue  # rendered as markdown table instead
            block_lines: list[str] = []
            for ln in b.get("lines", []):
                text = "".join(sp.get("text", "") for sp in ln.get("spans", [])).strip()
                if not text:
                    continue
                size = max(
                    (float(sp.get("size", 0)) for sp in ln.get("spans", []) if sp.get("text", "").strip()),
                    default=0.0,
                )
                bold = any("bold" in sp.get("font", "").lower() for sp in ln.get("spans", []))
                if (
                    median
                    and len(text) < 120
                    and (size >= median * 1.3 or (bold and size >= median * 1.1 and len(text) < 80))
                ):
                    level = 2 if size >= median * 1.5 else 3
                    block_lines.append("\n" + "#" * level + " " + text + "\n")
                else:
                    block_lines.append(text)
            if block_lines:
                lines_out.append("\n".join(block_lines))
    except Exception:  # noqa: BLE001
        lines_out = [page.get_text("text")]
    text = "\n\n".join(lines_out)
    if tables_md:
        text += "\n\n" + "\n\n".join(tables_md)
    return text.strip(), len(tables_md)


def _inside(inner: Any, outer: Any) -> bool:
    x0, y0, x1, y1 = inner
    ox0, oy0, ox1, oy1 = outer
    return x0 >= ox0 - 2 and y0 >= oy0 - 2 and x1 <= ox1 + 2 and y1 <= oy1 + 2


def parse_pdf_local(path: Path) -> ParsedDocument:
    import pymupdf

    doc = pymupdf.open(path)
    pages: list[str] = []
    total_chars = 0
    table_pages = 0
    image_count = 0
    for i, page in enumerate(doc):
        md, ntab = _pdf_page_markdown(page)
        total_chars += len(md)
        table_pages += 1 if ntab else 0
        try:
            image_count += len(page.get_images(full=False))
        except Exception:  # noqa: BLE001
            pass
        pages.append(f"<!-- page: {i + 1} -->\n{md}")
    n = max(len(doc), 1)
    meta = doc.metadata or {}
    title = (meta.get("title") or "").strip() or None
    markdown = "\n\n".join(pages)
    complexity = {
        "chars_per_page": total_chars / n,
        "table_page_ratio": table_pages / n,
        "images_per_page": image_count / n,
        "pages": len(doc),
    }
    return ParsedDocument(
        markdown=markdown,
        title=title or _title_from_text(markdown),
        page_count=len(doc),
        parser="pymupdf",
        metadata={"complexity": complexity, "pdf_metadata": {k: v for k, v in meta.items() if v}},
    )


def pdf_needs_llamaparse(c: dict[str, float]) -> bool:
    return c["chars_per_page"] < 200 or c["table_page_ratio"] > 0.2 or c["images_per_page"] > 1.5


def parse_pdf(path: Path, settings: Settings) -> ParsedDocument:
    local = parse_pdf_local(path)
    mode = settings.pdf_parser
    want_remote = mode == "llamaparse" or (
        mode == "auto" and pdf_needs_llamaparse(local.metadata["complexity"])
    )
    if want_remote and settings.llamaparse_api_key:
        from ..providers.parse import LlamaParseClient

        try:
            client = LlamaParseClient(
                settings.llamaparse_api_key,
                settings.llamaparse_base_url,
                tier=settings.llamaparse_tier,
                version=settings.llamaparse_version,
                language=settings.llamaparse_language or None,
            )
            res = client.parse(path, "application/pdf")
            return ParsedDocument(
                markdown=res.markdown,
                title=local.title or _title_from_text(res.markdown),
                page_count=len(res.pages) or local.page_count,
                parser=f"llamaparse:{settings.llamaparse_tier}",
                metadata={**local.metadata, **res.metadata},
            )
        except Exception as e:  # noqa: BLE001
            log.warning("LlamaParse failed for %s (%s); using local text", path.name, e)
            local.warnings.append(f"llamaparse_failed: {e}")
    elif want_remote:
        local.warnings.append("complex PDF but LLAMAPARSE_API_KEY not set; local extraction only")
    return local


# ---------------------------------------------------------------------------
# Office
# ---------------------------------------------------------------------------


def parse_docx(path: Path) -> ParsedDocument:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    out: list[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            p = Paragraph(child, doc)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            m = re.match(r"heading\s*(\d)", style, re.IGNORECASE)
            if m:
                out.append("\n" + "#" * min(int(m.group(1)), 6) + " " + text + "\n")
            elif style.lower().startswith("title"):
                out.append("\n# " + text + "\n")
            elif style.lower().startswith("list"):
                out.append("- " + text)
            else:
                out.append(text)
        elif child.tag == qn("w:tbl"):
            t = Table(child, doc)
            rows = [[c.text for c in r.cells] for r in t.rows]
            out.append("\n" + _md_table(rows) + "\n")
    markdown = "\n\n".join(out)
    title = (doc.core_properties.title or "").strip() or _title_from_text(markdown)
    created = doc.core_properties.created
    return ParsedDocument(
        markdown=markdown, title=title, parser="python-docx", doc_date=created.date() if created else None
    )


def parse_pptx(path: Path) -> ParsedDocument:
    from pptx import Presentation

    prs = Presentation(str(path))
    out: list[str] = []
    for n, slide in enumerate(prs.slides, start=1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        out.append(f"<!-- page: {n} -->\n## Slide {n}" + (f": {title}" if title else ""))
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[c.text for c in r.cells] for r in shape.table.rows]
                out.append(_md_table(rows))
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                if text.strip():
                    out.append(text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append("Speaker notes: " + notes)
    markdown = "\n\n".join(out)
    first = next((s for s in out if s.startswith("<!-- page: 1 -->")), "")
    title = (prs.core_properties.title or "").strip() or None
    if not title and "## Slide 1: " in first:
        title = first.split("## Slide 1: ", 1)[1].splitlines()[0].strip()
    return ParsedDocument(markdown=markdown, title=title, page_count=len(prs.slides), parser="python-pptx")


# ---------------------------------------------------------------------------
# Spreadsheets and databases
# ---------------------------------------------------------------------------


def parse_xlsx(path: Path) -> ParsedDocument:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    out: list[str] = []
    for ws in wb.worksheets:
        rows: list[list[str]] = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_SHEET_ROWS:
                out.append(f"_(sheet truncated at {MAX_SHEET_ROWS} rows)_")
                break
            cells = [_cell(v) for v in row]
            if any(c for c in cells):
                rows.append(cells)
        out.append(f"## Sheet: {ws.title}\n\n" + (_md_table(rows) if rows else "_(empty)_"))
    return ParsedDocument(
        markdown="\n\n".join(out), title=path.stem, source_type="spreadsheet", parser="openpyxl"
    )


def parse_csv(path: Path) -> ParsedDocument:
    raw = path.read_bytes().decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [[c for c in r] for r in csv.reader(raw.splitlines(), dialect) if any(c.strip() for c in r)]
    rows = rows[: MAX_SHEET_ROWS + 1]
    return ParsedDocument(
        markdown=f"## {path.stem}\n\n" + _md_table(rows),
        title=path.stem,
        source_type="spreadsheet",
        parser="csv",
    )


def parse_sqlite(path: Path) -> ParsedDocument:
    conn = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    try:
        out: list[str] = []
        for name, sql in conn.execute(
            "SELECT name, sql FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'"
        ):
            out.append(f"## Table: {name}\n\n```sql\n{sql}\n```")
            cur = conn.execute(f'SELECT * FROM "{name}" LIMIT {MAX_DB_ROWS}')
            cols = [d[0] for d in cur.description]
            rows = [[_cell(v) for v in r] for r in cur.fetchall()]
            out.append(_md_table([cols] + rows) if rows else "_(no rows)_")
    finally:
        conn.close()
    return ParsedDocument(markdown="\n\n".join(out), title=path.stem, source_type="database", parser="sqlite")


# ---------------------------------------------------------------------------
# E-mail
# ---------------------------------------------------------------------------


def _email_body(msg: email.message.Message) -> str:
    body = msg.get_body(preferencelist=("plain", "html")) if hasattr(msg, "get_body") else None
    if body is None:
        return ""
    content = body.get_content()
    if body.get_content_type() == "text/html":
        return html_to_text(content)
    return content.strip()


def _email_markdown(msg: Any, level: int = 1) -> tuple[str, date | None, str]:
    subject = str(msg.get("subject", "") or "(no subject)")
    sent: date | None = None
    try:
        dt = parsedate_to_datetime(msg.get("date")) if msg.get("date") else None
        sent = dt.date() if dt else None
    except Exception:  # noqa: BLE001
        sent = None
    header = [
        "#" * level + " " + subject,
        f"**From:** {msg.get('from', '')}",
        f"**To:** {msg.get('to', '')}",
    ]
    if msg.get("cc"):
        header.append(f"**Cc:** {msg.get('cc')}")
    header.append(f"**Date:** {msg.get('date', '')}")
    attachments = (
        [p.get_filename() for p in msg.iter_attachments()] if hasattr(msg, "iter_attachments") else []
    )
    attachments = [a for a in attachments if a]
    if attachments:
        header.append("**Attachments:** " + ", ".join(attachments))
    body = _email_body(msg)
    body = re.sub(r"\n{3,}", "\n\n", body)
    return "\n".join(header) + "\n\n" + body, sent, subject


def parse_eml(path: Path) -> ParsedDocument:
    msg = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
    md, sent, subject = _email_markdown(msg)
    return ParsedDocument(markdown=md, title=subject, source_type="email", doc_date=sent, parser="eml")


def parse_mbox(path: Path) -> ParsedDocument:
    box = mailbox.mbox(
        str(path), factory=lambda f: email.message_from_binary_file(f, policy=email.policy.default)
    )
    parts: list[str] = []
    dates: list[date] = []
    for i, msg in enumerate(box):
        if i >= 2000:
            parts.append("_(mailbox truncated at 2000 messages)_")
            break
        md, sent, _ = _email_markdown(msg, level=2)
        if sent:
            dates.append(sent)
        parts.append(md)
    return ParsedDocument(
        markdown=f"# Mailbox {path.stem}\n\n" + "\n\n---\n\n".join(parts),
        title=f"Mailbox {path.stem}",
        source_type="email",
        doc_date=max(dates) if dates else None,
        parser="mbox",
        metadata={"messages": len(parts)},
    )


# ---------------------------------------------------------------------------
# Chat exports (WhatsApp text, Telegram JSON)
# ---------------------------------------------------------------------------

_WA_RE = re.compile(
    r"^\[?(?P<d>\d{1,4}[./-]\d{1,2}[./-]\d{2,4}),?\s+(?P<t>\d{1,2}:\d{2}(?::\d{2})?\s?(?:[APap]\.?[Mm]\.?)?)\]?\s*[-–]?\s*(?P<who>[^:]{1,80}?):\s(?P<msg>.*)$"
)


def looks_like_whatsapp(text: str) -> bool:
    lines = [ln for ln in text.splitlines()[:60] if ln.strip()]
    hits = sum(1 for ln in lines if _WA_RE.match(normalize_persian(ln)))
    return len(lines) >= 3 and hits >= max(2, len(lines) // 3)


def _parse_wa_date(s: str, dayfirst: bool) -> date | None:
    nums = [int(x) for x in re.split(r"[./-]", s)]
    if len(nums) != 3:
        return None
    a, b, c = nums
    try:
        if a > 31:  # y-m-d
            return date(a, b, c)
        y = c if c > 99 else 2000 + c
        if y > 1500:  # Jalali years are stored as written; we cannot convert without context
            pass
        d, m = (a, b) if dayfirst else (b, a)
        return date(y, m, d)
    except ValueError:
        return None


def parse_whatsapp(text: str, title: str | None = None) -> ParsedDocument:
    text = normalize_persian(text)
    lines = text.splitlines()
    firsts, seconds = [], []
    for ln in lines:
        m = _WA_RE.match(ln)
        if m:
            parts = [int(x) for x in re.split(r"[./-]", m.group("d"))]
            if len(parts) == 3 and parts[0] <= 31:
                firsts.append(parts[0])
                seconds.append(parts[1])
    dayfirst = not (any(s > 12 for s in seconds) and not any(f > 12 for f in firsts))
    out: list[str] = []
    cur_day: date | None = None
    last_day: date | None = None
    participants: set[str] = set()
    for ln in lines:
        m = _WA_RE.match(ln)
        if not m:
            if out and ln.strip():
                out[-1] += "\n" + ln.strip()
            continue
        d = _parse_wa_date(m.group("d"), dayfirst)
        if d and d != cur_day:
            out.append(f"## {d.isoformat()}")
            cur_day = d
            last_day = d
        who = m.group("who").strip()
        participants.add(who)
        out.append(f"**{who}** ({m.group('t').strip()}): {m.group('msg').strip()}")
    md = (f"# {title}\n\n" if title else "") + "\n\n".join(out)
    return ParsedDocument(
        markdown=md,
        title=title,
        source_type="chat",
        doc_date=last_day,
        parser="whatsapp",
        metadata={
            "participants": sorted(participants),
            "messages": sum(1 for x in out if x.startswith("**")),
        },
    )


def looks_like_telegram(obj: Any) -> bool:
    return (
        isinstance(obj, dict) and isinstance(obj.get("messages"), list) and ("name" in obj or "type" in obj)
    )


def _tg_text(t: Any) -> str:
    if isinstance(t, str):
        return t
    if isinstance(t, list):
        return "".join(p if isinstance(p, str) else str(p.get("text", "")) for p in t)
    return ""


def parse_telegram(obj: dict[str, Any], fallback_title: str | None = None) -> ParsedDocument:
    title = obj.get("name") or fallback_title
    out: list[str] = []
    cur_day: str | None = None
    last_day: date | None = None
    participants: set[str] = set()
    n = 0
    for m in obj.get("messages", []):
        if m.get("type") not in (None, "message"):
            continue
        text = _tg_text(m.get("text", "")).strip()
        if not text and not m.get("file") and not m.get("photo"):
            continue
        when = str(m.get("date", ""))[:16].replace("T", " ")
        day = when[:10]
        if day and day != cur_day:
            out.append(f"## {day}")
            cur_day = day
            try:
                last_day = date.fromisoformat(day)
            except ValueError:
                pass
        who = str(m.get("from") or m.get("actor") or "unknown")
        participants.add(who)
        body = text or ("[media]" if m.get("photo") or m.get("file") else "")
        out.append(f"**{who}** ({when[11:]}): {normalize_persian(body)}")
        n += 1
    md = (f"# {title}\n\n" if title else "") + "\n\n".join(out)
    return ParsedDocument(
        markdown=md,
        title=title,
        source_type="chat",
        doc_date=last_day,
        parser="telegram",
        metadata={"participants": sorted(participants), "messages": n},
    )


# ---------------------------------------------------------------------------
# dispatcher
# ---------------------------------------------------------------------------


def parse_file(path: Path, settings: Settings, filename: str | None = None) -> ParsedDocument:
    name = filename or path.name
    ext = Path(name).suffix.lower()
    if ext == ".pdf":
        doc = parse_pdf(path, settings)
    elif ext == ".docx":
        doc = parse_docx(path)
    elif ext == ".pptx":
        doc = parse_pptx(path)
    elif ext in (".xlsx", ".xlsm"):
        doc = parse_xlsx(path)
    elif ext in (".csv", ".tsv"):
        doc = parse_csv(path)
    elif ext in (".sqlite", ".sqlite3", ".db"):
        doc = parse_sqlite(path)
    elif ext == ".eml":
        doc = parse_eml(path)
    elif ext == ".mbox":
        doc = parse_mbox(path)
    elif ext in (".html", ".htm"):
        raw = path.read_bytes().decode("utf-8", errors="replace")
        text = html_to_text(raw)
        m = re.search(r"<title>(.*?)</title>", raw, re.I | re.S)
        doc = ParsedDocument(
            markdown=text,
            title=html.unescape(m.group(1).strip()) if m else _title_from_text(text),
            parser="html",
            source_type="web",
        )
    elif ext == ".json":
        raw = path.read_bytes().decode("utf-8", errors="replace")
        try:
            obj = json.loads(raw)
        except json.JSONDecodeError as e:
            raise UnsupportedFormat(f"invalid JSON: {e}") from e
        if looks_like_telegram(obj):
            doc = parse_telegram(obj, Path(name).stem)
        else:
            doc = ParsedDocument(
                markdown=f"# {Path(name).stem}\n\n```json\n{json.dumps(obj, ensure_ascii=False, indent=1)[:400_000]}\n```",
                title=Path(name).stem,
                parser="json",
            )
    elif ext in TEXT_EXTS or ext == "":
        try:
            raw = path.read_bytes().decode("utf-8-sig")
        except UnicodeDecodeError as e:
            raise UnsupportedFormat(f"{name}: not UTF-8 text") from e
        if looks_like_whatsapp(raw):
            doc = parse_whatsapp(raw, Path(name).stem)
        else:
            doc = ParsedDocument(
                markdown=raw,
                title=_title_from_text(raw) or Path(name).stem,
                parser="text",
                source_type="note",
            )
    else:
        raise UnsupportedFormat(f"unsupported file type: {ext or name}")

    doc.markdown = normalize_persian(doc.markdown).strip()
    if doc.title:
        doc.title = normalize_persian(doc.title).strip()[:300]
    doc.lang = doc.lang or detect_lang(doc.markdown[:20000])
    return doc
