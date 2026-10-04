from __future__ import annotations

import email.message
from pathlib import Path

from mybrain.config import Settings
from mybrain.ingest.parsers import (
    html_to_text,
    looks_like_whatsapp,
    parse_file,
    parse_telegram,
    parse_whatsapp,
)

SETTINGS = Settings(pdf_parser="local", llamaparse_api_key="")

WA = """12/31/23, 10:15 PM - Ali Rezaei: Salam, did you see the term sheet?
12/31/23, 10:16 PM - Sara: Yes, the valuation is 40m
This continues the previous line.
1/1/24, 9:00 AM - Ali Rezaei: ok let's meet Monday
"""


def test_whatsapp_detection_and_parsing():
    assert looks_like_whatsapp(WA)
    doc = parse_whatsapp(WA, "deal chat")
    assert doc.source_type == "chat"
    assert "## 2023-12-31" in doc.markdown and "## 2024-01-01" in doc.markdown
    assert "**Sara** (10:16 PM): Yes, the valuation is 40m\nThis continues the previous line." in doc.markdown
    assert set(doc.metadata["participants"]) == {"Ali Rezaei", "Sara"}
    assert doc.doc_date is not None and doc.doc_date.isoformat() == "2024-01-01"


def test_whatsapp_dayfirst_bracket_format():
    text = "[31/12/2023, 22:15:03] علی رضایی: سلام\n[01/01/2024, 09:00:00] Sara: hi\n[02/01/2024, 09:00:00] Sara: again\n"
    assert looks_like_whatsapp(text)
    doc = parse_whatsapp(text)
    assert "## 2023-12-31" in doc.markdown and "## 2024-01-02" in doc.markdown


def test_telegram_export():
    obj = {
        "name": "Deal group",
        "type": "private_supergroup",
        "messages": [
            {"id": 1, "type": "message", "date": "2024-02-03T10:00:00", "from": "Sara", "text": "hello"},
            {"id": 2, "type": "service", "date": "2024-02-03T10:01:00", "actor": "x", "text": ""},
            {
                "id": 3,
                "type": "message",
                "date": "2024-02-04T11:00:00",
                "from": "Ali",
                "text": ["see ", {"type": "link", "text": "http://x"}],
            },
        ],
    }
    doc = parse_telegram(obj)
    assert doc.title == "Deal group" and doc.metadata["messages"] == 2
    assert "**Ali** (11:00): see http://x" in doc.markdown


def test_html_to_text_headings():
    text = html_to_text(
        "<html><head><title>T</title><style>x{}</style></head><body><h1>Head</h1><p>Para &amp; more</p><script>bad()</script></body></html>"
    )
    assert "# Head" in text and "Para & more" in text and "bad()" not in text


def test_parse_file_dispatch(tmp_path: Path):
    md = tmp_path / "note.md"
    md.write_text(
        "# Meeting with Pars Holding\n\nDiscussed valuation with Ali Rezaei (علي رضايي).\n", encoding="utf-8"
    )
    doc = parse_file(md, SETTINGS)
    assert doc.title == "Meeting with Pars Holding" and doc.parser == "text"
    assert "علی رضایی" in doc.markdown  # Arabic letters normalized to Persian

    csvf = tmp_path / "table.csv"
    csvf.write_text("name,amount\nPars Holding,100\nTehran Metro,200\n", encoding="utf-8")
    doc = parse_file(csvf, SETTINGS)
    assert doc.source_type == "spreadsheet" and "| Pars Holding | 100 |" in doc.markdown

    msg = email.message.EmailMessage()
    msg["From"] = "ali@pars.example"
    msg["To"] = "me@example.com"
    msg["Subject"] = "Term sheet"
    msg["Date"] = "Mon, 01 Jan 2024 10:00:00 +0330"
    msg.set_content("Please find the term sheet attached.")
    eml = tmp_path / "mail.eml"
    eml.write_bytes(bytes(msg))
    doc = parse_file(eml, SETTINGS)
    assert (
        doc.source_type == "email" and doc.title == "Term sheet" and doc.doc_date.isoformat() == "2024-01-01"
    )
    assert "**From:** ali@pars.example" in doc.markdown and "term sheet attached" in doc.markdown

    wa = tmp_path / "chat.txt"
    wa.write_text(WA, encoding="utf-8")
    assert parse_file(wa, SETTINGS).parser == "whatsapp"


def test_parse_office_and_pdf(tmp_path: Path):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    d = Document()
    d.add_heading("Investment memo", level=1)
    d.add_paragraph("Pars Holding seeks 40m.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Year", "Revenue"
    t.cell(1, 0).text, t.cell(1, 1).text = "2023", "130"
    p = tmp_path / "memo.docx"
    d.save(str(p))
    doc = parse_file(p, SETTINGS)
    assert "# Investment memo" in doc.markdown and "| 2023 | 130 |" in doc.markdown

    wb = Workbook()
    ws = wb.active
    ws.title = "Cap table"
    ws.append(["Holder", "Share"])
    ws.append(["Ali Rezaei", 0.25])
    x = tmp_path / "cap.xlsx"
    wb.save(str(x))
    doc = parse_file(x, SETTINGS)
    assert "## Sheet: Cap table" in doc.markdown and "| Ali Rezaei | 0.25 |" in doc.markdown

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Deal overview"
    slide.placeholders[1].text = "Valuation 40m"
    pp = tmp_path / "deck.pptx"
    prs.save(str(pp))
    doc = parse_file(pp, SETTINGS)
    assert (
        "## Slide 1: Deal overview" in doc.markdown
        and "Valuation 40m" in doc.markdown
        and doc.page_count == 1
    )

    import pymupdf

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Pars Holding annual report", fontsize=20)
    page.insert_text((72, 120), "Revenue grew to 130 in 2023. The CFO is Ali Rezaei.", fontsize=11)
    pf = tmp_path / "report.pdf"
    pdf.save(str(pf))
    doc = parse_file(pf, SETTINGS)
    assert doc.parser == "pymupdf" and doc.page_count == 1
    assert "Pars Holding annual report" in doc.markdown and "Ali Rezaei" in doc.markdown
    assert "<!-- page: 1 -->" in doc.markdown
